from __future__ import annotations

import csv
import math
import re
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "report_assets"
SOURCE_MD = DOCS / "bao_cao_chuong_1_4_hoan_chinh.md"
OUTPUT_DOCX = DOCS / "Bao_cao_TTTN_AI_Generated_Image_Detection.docx"
METRICS_CSV = ROOT / "benchmark" / "results" / "openfake_1k" / "metrics.csv"
REAL_DIR = ROOT / "benchmark" / "data" / "openfake_1k" / "images" / "real"
FAKE_DIR = ROOT / "benchmark" / "data" / "openfake_1k" / "images" / "fake"
PRED_DIR = ROOT / "benchmark" / "results" / "openfake_1k" / "predictions"


TABLE_CAPTIONS = [
    "Bảng 1.1. Công nghệ sử dụng trong đề tài",
    "Bảng 2.1. Tác nhân và vai trò trong hệ thống",
    "Bảng 2.2. Ánh xạ Label_A trong Defactify Image Dataset",
    "Bảng 2.3. Ánh xạ Label_B trong Defactify Image Dataset",
    "Bảng 2.4. Hai model ONNX tích hợp trong web app",
    "Bảng 2.5. Các thành phần trong hướng nghiên cứu SFW-SwinCBM",
    "Bảng 2.6. Các API backend chính của hệ thống",
    "Bảng 2.7. Các biến môi trường cấu hình hệ thống",
    "Bảng 2.8. Các trường hợp lỗi cần xử lý",
    "Bảng 3.1. Kết quả benchmark OpenFake 500",
    "Bảng 3.2. Ma trận nhầm lẫn tại threshold 0.5",
    "Bảng 3.3. Các nhóm kiểm thử trong dự án",
]


FIGURE_HOOKS = {
    "## 2.2. Use case tong quan": [
        ("fig_2_1_usecase.png", "Hình 2.1. Luồng use case tổng quan của hệ thống")
    ],
    "## 2.3. Du lieu su dung": [
        ("fig_2_2_dataset_samples.png", "Hình 2.2. Ví dụ ảnh Real và Fake trong bộ dữ liệu benchmark OpenFake")
    ],
    "## 2.4. Kien truc tong the cua he thong": [
        ("fig_2_3_architecture.png", "Hình 2.3. Kiến trúc tổng thể của hệ thống web nhận diện ảnh AI")
    ],
    "## 2.5. Thiet ke mo hinh nhan dien": [
        ("fig_2_4_model_pipeline.png", "Hình 2.4. Pipeline suy luận và chọn kết quả cuối cùng từ hai model ONNX")
    ],
    "## 3.3. Trien khai giao dien web": [
        ("fig_3_1_ui_mock.png", "Hình 3.1. Bố cục giao diện web của hệ thống")
    ],
    "### 3.7.3. Ket qua benchmark": [
        ("fig_3_2_benchmark_metrics.png", "Hình 3.2. So sánh Accuracy, F1 fake và ROC-AUC giữa các model"),
        ("fig_3_3_score_distribution.png", "Hình 3.3. Phân bố fake score của các model trên OpenFake 500"),
    ],
    "### 3.7.5. Ma tran nham lan o threshold 0.5": [
        ("fig_3_4_confusion_matrices.png", "Hình 3.4. Ma trận nhầm lẫn của ba model tại threshold 0.5")
    ],
}


def font_path(*names: str) -> str | None:
    for name in names:
        path = Path("C:/Windows/Fonts") / name
        if path.exists():
            return str(path)
    return None


FONT_REG = font_path("times.ttf", "arial.ttf")
FONT_BOLD = font_path("timesbd.ttf", "arialbd.ttf")


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold and FONT_BOLD else FONT_REG
    if path:
        return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def ensure_assets() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    create_usecase_diagram(ASSETS / "fig_2_1_usecase.png")
    create_dataset_collage(ASSETS / "fig_2_2_dataset_samples.png")
    create_architecture_diagram(ASSETS / "fig_2_3_architecture.png")
    create_model_pipeline(ASSETS / "fig_2_4_model_pipeline.png")
    create_ui_mock(ASSETS / "fig_3_1_ui_mock.png")
    create_benchmark_chart(ASSETS / "fig_3_2_benchmark_metrics.png")
    create_score_distribution(ASSETS / "fig_3_3_score_distribution.png")
    create_confusion_matrices(ASSETS / "fig_3_4_confusion_matrices.png")


def draw_centered(draw: ImageDraw.ImageDraw, xyxy, text: str, font, fill=(20, 31, 46), spacing=4) -> None:
    x1, y1, x2, y2 = xyxy
    max_width = x2 - x1 - 24
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    total_h = sum(draw.textbbox((0, 0), line, font=font)[3] for line in lines) + spacing * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += (bbox[3] - bbox[1]) + spacing


def round_box(draw, xyxy, text, fill, outline, font, radius=18, text_fill=(20, 31, 46)) -> None:
    draw.rounded_rectangle(xyxy, radius=radius, fill=fill, outline=outline, width=3)
    draw_centered(draw, xyxy, text, font, fill=text_fill)


def arrow(draw, start, end, fill=(47, 79, 112), width=4) -> None:
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    angle = math.atan2(ey - sy, ex - sx)
    size = 13
    pts = [
        (ex, ey),
        (ex - size * math.cos(angle - 0.45), ey - size * math.sin(angle - 0.45)),
        (ex - size * math.cos(angle + 0.45), ey - size * math.sin(angle + 0.45)),
    ]
    draw.polygon(pts, fill=fill)


def create_usecase_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#f7fafc")
    d = ImageDraw.Draw(img)
    title = load_font(42, True)
    box_font = load_font(28, True)
    small = load_font(24)
    d.text((60, 40), "Luồng sử dụng hệ thống AI Generated Image Detection", font=title, fill="#111827")
    boxes = [
        ((70, 210, 350, 350), "Người dùng\nmở web"),
        ((430, 210, 710, 350), "Kiểm tra\nMODEL ON/OFF"),
        ((790, 125, 1070, 265), "Upload file\nảnh"),
        ((790, 315, 1070, 455), "Nhập URL và\nLoad URL"),
        ((1150, 210, 1430, 350), "Phân tích\nảnh"),
        ((1150, 530, 1430, 720), "Hiển thị Real/AI,\nđiểm AI,\nvoting model"),
    ]
    for xy, text in boxes:
        round_box(d, xy, text, "#e0f2fe", "#0284c7", box_font)
    arrow(d, (350, 280), (430, 280))
    arrow(d, (710, 260), (790, 205))
    arrow(d, (710, 300), (790, 385))
    arrow(d, (1070, 205), (1150, 260))
    arrow(d, (1070, 385), (1150, 300))
    arrow(d, (1290, 350), (1290, 530))
    d.text((435, 385), "Nếu MODEL OFF: khóa input và không cho phân tích", font=small, fill="#b45309")
    d.text((435, 420), "Nếu MODEL ON: cho phép chuẩn bị ảnh đầu vào", font=small, fill="#047857")
    img.save(path, quality=95)


def create_dataset_collage(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    title = load_font(42, True)
    label_font = load_font(32, True)
    d.text((60, 35), "Mẫu ảnh benchmark OpenFake: Real và Fake", font=title, fill="#111827")
    real_files = sorted(REAL_DIR.glob("*"))[:4]
    fake_files = sorted(FAKE_DIR.glob("*"))[:4]
    cells = [(80 + i * 360, 160, 380 + i * 360, 460) for i in range(4)]
    cells += [(80 + i * 360, 545, 380 + i * 360, 845) for i in range(4)]
    for idx, file in enumerate(real_files + fake_files):
        with Image.open(file) as sample:
            sample = sample.convert("RGB")
            sample.thumbnail((300, 300), Image.Resampling.LANCZOS)
            x1, y1, x2, y2 = cells[idx]
            d.rounded_rectangle((x1 - 8, y1 - 8, x2 + 8, y2 + 8), 14, fill="#ffffff", outline="#cbd5e1", width=2)
            px = x1 + (300 - sample.width) // 2
            py = y1 + (300 - sample.height) // 2
            img.paste(sample, (px, py))
    d.text((80, 105), "REAL", font=label_font, fill="#047857")
    d.text((80, 490), "FAKE / AI-GENERATED", font=label_font, fill="#b91c1c")
    img.save(path, quality=95)


def create_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    title = load_font(42, True)
    box_font = load_font(27, True)
    small = load_font(22)
    d.text((60, 40), "Kiến trúc tổng thể hệ thống", font=title, fill="#111827")
    boxes = [
        ((70, 245, 330, 410), "User\nBrowser"),
        ((430, 190, 760, 465), "Web UI\nHTML/CSS/JavaScript\nUpload, URL, Preview"),
        ((860, 165, 1190, 490), "Backend Flask API\nValidate input\nModel state\nJSON response"),
        ((1280, 190, 1530, 465), "AI Worker\nONNX Runtime\n2 model sessions"),
        ((860, 600, 1190, 760), "Model files\nhybrid_xrayon_physical.onnx\nxrayon_rgb_only.onnx"),
    ]
    colors = ["#dcfce7", "#e0f2fe", "#ede9fe", "#fee2e2", "#fef3c7"]
    outlines = ["#16a34a", "#0284c7", "#7c3aed", "#dc2626", "#d97706"]
    for (xy, text), fill, out in zip(boxes, colors, outlines):
        round_box(d, xy, text, fill, out, box_font)
    arrow(d, (330, 330), (430, 330))
    arrow(d, (760, 330), (860, 330))
    arrow(d, (1190, 330), (1280, 330))
    arrow(d, (1190, 675), (1280, 440))
    d.text((450, 500), "GET /api/model-state, POST /api/load-url, POST /api/analyze", font=small, fill="#334155")
    d.text((900, 525), "Worker process tách riêng để giải phóng RAM khi MODEL OFF", font=small, fill="#334155")
    img.save(path, quality=95)


def create_model_pipeline(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    title = load_font(42, True)
    box_font = load_font(25, True)
    d.text((60, 40), "Pipeline suy luận và chọn kết quả cuối cùng", font=title, fill="#111827")
    round_box(d, (80, 355, 300, 505), "Ảnh RGB\nUpload/URL", "#dcfce7", "#16a34a", box_font)
    round_box(d, (390, 250, 700, 400), "Hybrid XRayon Physical\nInput 224x224\nONNX", "#fee2e2", "#dc2626", box_font)
    round_box(d, (390, 510, 700, 660), "XRayon RGB Only\nInput 256x256\nONNX", "#e0f2fe", "#0284c7", box_font)
    round_box(d, (790, 250, 1080, 400), "Output model 1\nprob_real, prob_ai\nconfidence", "#fff7ed", "#ea580c", box_font)
    round_box(d, (790, 510, 1080, 660), "Output model 2\nprob_real, prob_ai\nconfidence", "#fff7ed", "#ea580c", box_font)
    round_box(d, (1180, 380, 1500, 550), "Chọn model có\nconfidence cao nhất\n=> final_label", "#ede9fe", "#7c3aed", box_font)
    for y in (325, 585):
        arrow(d, (300, 430), (390, y))
        arrow(d, (700, y), (790, y))
        arrow(d, (1080, y), (1180, 465))
    img.save(path, quality=95)


def create_ui_mock(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#d6e0eb")
    d = ImageDraw.Draw(img)
    title = load_font(36, True)
    h = load_font(28, True)
    p = load_font(23)
    d.rounded_rectangle((80, 70, 1520, 830), 24, fill="#f1f5fa", outline="#8da0b5", width=3)
    d.rectangle((80, 70, 1520, 130), fill="#dbe5f0")
    d.ellipse((110, 92, 132, 114), fill="#ff453a")
    d.ellipse((145, 92, 167, 114), fill="#ff9f0a")
    d.ellipse((180, 92, 202, 114), fill="#30d158")
    d.text((230, 90), "~/ai-image-detection", font=p, fill="#253447")
    d.text((120, 165), "AI Generated Image Detection", font=title, fill="#0b1220")
    d.text((120, 215), "Kiểm tra ảnh thật hay ảnh do AI tạo", font=p, fill="#253447")
    d.rounded_rectangle((120, 285, 500, 770), 12, fill="#e8eff7", outline="#8da0b5", width=2)
    d.text((150, 315), "Ảnh đầu vào", font=h, fill="#0b1220")
    d.rounded_rectangle((150, 375, 470, 530), 10, fill="#dbe5ef", outline="#8da0b5", width=2)
    d.text((205, 430), "Chọn ảnh / kéo thả", font=p, fill="#253447")
    d.rounded_rectangle((150, 580, 350, 625), 8, fill="#dbe5ef", outline="#8da0b5", width=2)
    d.text((165, 592), "URL ảnh", font=p, fill="#253447")
    d.rounded_rectangle((365, 580, 470, 625), 8, fill="#005eb8")
    d.text((388, 592), "Load", font=p, fill="#ffffff")
    d.rounded_rectangle((150, 690, 470, 740), 8, fill="#005eb8")
    d.text((260, 704), "Phân tích", font=p, fill="#ffffff")
    d.rounded_rectangle((560, 285, 1480, 770), 12, fill="#e8eff7", outline="#8da0b5", width=2)
    d.text((590, 315), "Kết quả", font=h, fill="#0b1220")
    d.rounded_rectangle((590, 370, 1450, 650), 10, fill="#e5edf6", outline="#8da0b5", width=2)
    d.text((945, 485), "Preview ảnh", font=h, fill="#4c6075")
    cards = [(590, 680, 850, 745), (890, 680, 1150, 745), (1190, 680, 1450, 745)]
    labels = ["Nhãn: Anh AI", "Điểm AI: 82%", "Xử lý: 214 ms"]
    for xy, label in zip(cards, labels):
        d.rounded_rectangle(xy, 8, fill="#dfe8f2", outline="#8da0b5", width=2)
        d.text((xy[0] + 18, xy[1] + 18), label, font=p, fill="#0b1220")
    img.save(path, quality=95)


def create_benchmark_chart(path: Path) -> None:
    df = pd.read_csv(METRICS_CSV)
    df = df[df["threshold"] == 0.5].copy()
    models = df["model"].tolist()
    x = np.arange(len(models))
    width = 0.23
    plt.figure(figsize=(12, 6), dpi=180)
    plt.bar(x - width, df["accuracy"], width, label="Accuracy", color="#2563eb")
    plt.bar(x, df["f1_fake"], width, label="F1 fake", color="#16a34a")
    plt.bar(x + width, df["roc_auc"], width, label="ROC-AUC", color="#dc2626")
    plt.xticks(x, models, rotation=10, ha="right")
    plt.ylim(0, 1)
    plt.ylabel("Giá trị metric")
    plt.title("So sánh benchmark OpenFake 500 tại threshold 0.5")
    plt.grid(axis="y", alpha=0.25)
    plt.legend()
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()


def create_score_distribution(path: Path) -> None:
    files = [
        ("Hybrid XRayon Physical", PRED_DIR / "raw_hybrid_xrayon_physical.csv"),
        ("XRayon RGB Only", PRED_DIR / "raw_xrayon_rgb_only.csv"),
        ("UniversalFakeDetect", PRED_DIR / "raw_universalfakedetect.csv"),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(14, 4.2), dpi=180, sharey=True)
    for ax, (name, file) in zip(axes, files):
        df = pd.read_csv(file)
        label_col = "label" if "label" in df.columns else ("y_true" if "y_true" in df.columns else None)
        score_col = "fake_score" if "fake_score" in df.columns else ("score" if "score" in df.columns else None)
        if label_col is None or score_col is None:
            # Fallback for current benchmark format.
            label_col = [c for c in df.columns if "label" in c.lower() or "true" in c.lower()][0]
            score_col = [c for c in df.columns if "score" in c.lower() or "fake" in c.lower()][-1]
        labels = df[label_col].astype(str).str.lower()
        real = df[labels.str.contains("real") | (labels == "0")][score_col]
        fake = df[labels.str.contains("fake") | labels.str.contains("ai") | (labels == "1")][score_col]
        ax.hist(real, bins=20, alpha=0.65, label="Real", color="#2563eb")
        ax.hist(fake, bins=20, alpha=0.65, label="Fake", color="#dc2626")
        ax.set_title(name)
        ax.set_xlabel("Fake score")
        ax.grid(axis="y", alpha=0.2)
    axes[0].set_ylabel("Số ảnh")
    axes[0].legend()
    fig.suptitle("Phân bố fake score trên OpenFake 500", y=1.02)
    plt.tight_layout()
    plt.savefig(path, bbox_inches="tight")
    plt.close()


def create_confusion_matrices(path: Path) -> None:
    df = pd.read_csv(METRICS_CSV)
    df = df[df["threshold"] == 0.5].copy()
    fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.3), dpi=180)
    for ax, (_, row) in zip(axes, df.iterrows()):
        matrix = np.array(
            [
                [row["true_real_pred_real"], row["true_real_pred_fake"]],
                [row["true_fake_pred_real"], row["true_fake_pred_fake"]],
            ],
            dtype=float,
        )
        im = ax.imshow(matrix, cmap="Blues")
        ax.set_title(row["model"], fontsize=10)
        ax.set_xticks([0, 1], labels=["Pred Real", "Pred Fake"])
        ax.set_yticks([0, 1], labels=["True Real", "True Fake"])
        for i in range(2):
            for j in range(2):
                ax.text(j, i, int(matrix[i, j]), ha="center", va="center", color="#111827", fontsize=12)
    fig.colorbar(im, ax=axes.ravel().tolist(), shrink=0.75)
    fig.suptitle("Ma trận nhầm lẫn tại threshold 0.5", y=1.02)
    plt.savefig(path, bbox_inches="tight")
    plt.close()


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def add_toc_field(paragraph, instr_text: str, placeholder: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(3)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15

    for i in range(1, 7):
        style = styles[f"Heading {i}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(13)
    styles["Heading 4"].font.size = Pt(13)
    styles["Heading 5"].font.size = Pt(12)
    styles["Heading 6"].font.size = Pt(12)

    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Báo cáo TTTN Đại học - Hệ thống nhận diện ảnh do AI tạo ra")
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(10)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.add_run("Nhóm_AIGID - Trang ")
        add_page_field(fp)


def add_paragraph(doc: Document, text: str, style: str | None = None, align=None, bold=False) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if style is None:
        r.font.size = Pt(13)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_caption(doc: Document, text: str, is_table: bool) -> None:
    style = "Heading 5" if is_table else "Heading 6"
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(12)


def add_image(doc: Document, filename: str, caption: str) -> None:
    path = ASSETS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.3))
    add_caption(doc, caption, is_table=False)


def parse_markdown_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    header = [c.strip() for c in lines[0].strip().strip("|").split("|")]
    rows: list[list[str]] = []
    for line in lines[2:]:
        rows.append([c.strip() for c in line.strip().strip("|").split("|")])
    return header, rows


def add_word_table(doc: Document, header: list[str], rows: list[list[str]], caption: str) -> None:
    add_caption(doc, caption, is_table=True)
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, col in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_text(cell, col, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_text(cells[idx], value)


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(10)


def render_markdown_body(doc: Document) -> None:
    text = SOURCE_MD.read_text(encoding="utf-8")
    lines = text.splitlines()
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    table_index = 0

    def flush_table():
        nonlocal table_lines, table_index
        if table_lines:
            header, rows = parse_markdown_table(table_lines)
            caption = TABLE_CAPTIONS[table_index] if table_index < len(TABLE_CAPTIONS) else f"Bảng {table_index + 1}. Bảng dữ liệu"
            table_index += 1
            add_word_table(doc, header, rows, caption)
            table_lines = []

    def flush_code():
        nonlocal code_lines
        if code_lines:
            add_code_block(doc, code_lines)
            code_lines = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("| "):
            table_lines.append(line)
            continue
        flush_table()

        if not line.strip() or line.strip() == "---":
            continue
        if line.startswith("# NOI DUNG") or line.startswith("Ghi chu bien tap") or line.startswith("De tai de xuat"):
            continue

        if line in FIGURE_HOOKS:
            for filename, caption in FIGURE_HOOKS[line]:
                add_image(doc, filename, caption)

        if line.startswith("# "):
            add_heading(doc, normalize_heading_text(line[2:]), 1)
        elif line.startswith("## "):
            add_heading(doc, normalize_heading_text(line[3:]), 2)
        elif line.startswith("### "):
            add_heading(doc, normalize_heading_text(line[4:]), 3)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.7)
            add_inline_runs(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.7)
            add_inline_runs(p, re.sub(r"^\d+\. ", "", line))
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, line)

    flush_table()
    flush_code()


def normalize_heading_text(text: str) -> str:
    text = text.strip()
    text = text.replace("CHUONG", "CHƯƠNG")
    text = text.replace("TONG QUAN DE TAI", "TỔNG QUAN ĐỀ TÀI")
    text = text.replace("PHAN TICH VA THIET KE HE THONG", "PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")
    text = text.replace("TRIEN KHAI VA DANH GIA HE THONG", "TRIỂN KHAI VÀ ĐÁNH GIÁ HỆ THỐNG")
    text = text.replace("KET LUAN VA HUONG PHAT TRIEN", "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    return text


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        mono = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else (part[1:-1] if mono else part)
        run = paragraph.add_run(clean)
        run.bold = bold
        run.font.name = "Consolas" if mono else "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas" if mono else "Times New Roman")
        run.font.size = Pt(11 if mono else 13)


def add_cover(doc: Document) -> None:
    for text, size, bold in [
        ("BỘ KHOA HỌC VÀ CÔNG NGHỆ", 13, True),
        ("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", 13, True),
        ("CƠ SỞ TẠI THÀNH PHỐ HỒ CHÍ MINH", 13, True),
        ("------------------------------", 13, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(size)
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO THỰC TẬP\nTỐT NGHIỆP ĐẠI HỌC")
    r.bold = True
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(30)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Đề tài: "Xây dựng hệ thống nhận diện ảnh do trí tuệ nhân tạo tạo ra"')
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for _ in range(3):
        doc.add_paragraph()
    info = [
        "Người hướng dẫn: TS. Huỳnh Trung Trụ",
        "Sinh viên 1: Trần Công Hậu        MSSV: N22DCCI010        Vai trò: Nhóm trưởng",
        "Sinh viên 2: Vũ Phạm Minh Thức    MSSV: N22DCVT099        Vai trò: Thành viên",
        "Sinh viên 3: Nguyễn Minh Quân     MSSV: N22DCDK071        Vai trò: Thành viên",
        "Ngành: Công nghệ thông tin",
    ]
    for item in info:
        add_paragraph(doc, item, bold=True)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TP. Hồ Chí Minh, tháng 08/2026")
    r.bold = True
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(13)
    doc.add_page_break()


def add_preliminary_pages(doc: Document) -> None:
    add_heading(doc, "LỜI CẢM ƠN", 1)
    thanks = (
        "Nhóm chúng em xin chân thành cảm ơn Học viện Công nghệ Bưu chính Viễn thông, "
        "các thầy cô trong Khoa Công nghệ thông tin đã tạo điều kiện để nhóm thực hiện đề tài thực tập tốt nghiệp. "
        "Đặc biệt, nhóm xin gửi lời cảm ơn đến TS. Huỳnh Trung Trụ đã hướng dẫn, góp ý và định hướng trong quá trình "
        "nghiên cứu, xây dựng hệ thống nhận diện ảnh do trí tuệ nhân tạo tạo ra. Những góp ý của thầy giúp nhóm hoàn thiện "
        "hơn cả phần kỹ thuật, cách trình bày báo cáo và phương pháp đánh giá kết quả."
    )
    add_paragraph(doc, thanks)
    add_paragraph(
        doc,
        "Trong quá trình thực hiện, do thời gian và kinh nghiệm còn hạn chế, báo cáo khó tránh khỏi thiếu sót. "
        "Nhóm mong nhận được ý kiến đóng góp của quý thầy cô để tiếp tục cải thiện hệ thống trong các phiên bản sau."
    )
    doc.add_page_break()

    add_heading(doc, "DANH MỤC CÁC KÝ HIỆU VÀ CHỮ VIẾT TẮT", 1)
    abbrev = [
        ("AI", "Artificial Intelligence - Trí tuệ nhân tạo"),
        ("XAI", "Explainable Artificial Intelligence - Trí tuệ nhân tạo có khả năng giải thích"),
        ("ONNX", "Open Neural Network Exchange"),
        ("API", "Application Programming Interface"),
        ("RGB", "Red Green Blue"),
        ("DCT", "Discrete Cosine Transform"),
        ("DWT", "Discrete Wavelet Transform"),
        ("UI", "User Interface"),
        ("URL", "Uniform Resource Locator"),
    ]
    add_word_table(doc, ["Từ viết tắt", "Diễn giải"], abbrev, "Bảng 0.1. Danh mục các ký hiệu và chữ viết tắt")
    doc.add_page_break()

    add_heading(doc, "MỤC LỤC", 1)
    p = doc.add_paragraph()
    add_toc_field(p, r'TOC \o "1-4" \h \z \u', "Bấm Ctrl+A rồi F9 trong Word để cập nhật mục lục.")
    doc.add_page_break()

    add_heading(doc, "DANH MỤC CÁC BẢNG", 1)
    p = doc.add_paragraph()
    add_toc_field(p, r'TOC \t "Heading 5,1" \h \z', "Bấm Ctrl+A rồi F9 trong Word để cập nhật danh mục bảng.")
    doc.add_page_break()

    add_heading(doc, "DANH MỤC CÁC HÌNH VẼ", 1)
    p = doc.add_paragraph()
    add_toc_field(p, r'TOC \t "Heading 6,1" \h \z', "Bấm Ctrl+A rồi F9 trong Word để cập nhật danh mục hình.")
    doc.add_page_break()

    add_heading(doc, "MỞ ĐẦU", 1)
    add_paragraph(
        doc,
        "Sự phát triển nhanh của các mô hình tạo ảnh bằng trí tuệ nhân tạo đã làm cho ảnh tổng hợp ngày càng khó phân biệt "
        "với ảnh chụp thật. Điều này tạo ra nhu cầu xây dựng các công cụ hỗ trợ kiểm tra nguồn gốc hình ảnh, đặc biệt trong "
        "bối cảnh thông tin số, truyền thông trực tuyến và mạng xã hội. Đề tài này tập trung xây dựng một hệ thống web local "
        "cho phép người dùng upload ảnh hoặc nhập URL, sau đó sử dụng các mô hình ONNX để dự đoán ảnh thật hay ảnh do AI tạo ra."
    )
    add_paragraph(
        doc,
        "Báo cáo trình bày từ cơ sở lý thuyết, phân tích yêu cầu, thiết kế hệ thống, triển khai web app, inference model, "
        "benchmark OpenFake 500 và các hướng phát triển tiếp theo. Nội dung được xây dựng dựa trên mã nguồn hiện có của dự án "
        "và kết quả benchmark đã chạy trong repository."
    )
    doc.add_page_break()


def main() -> None:
    ensure_assets()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_preliminary_pages(doc)
    render_markdown_body(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
